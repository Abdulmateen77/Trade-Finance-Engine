"""
Contract .docx generator for OceanX pipeline.

Converts the markdown contract draft from the Contract Agent into a
formatted Word document using python-docx.
"""

import os
import re
from datetime import date

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _slugify(name: str) -> str:
    """Lowercase, replace non-alphanumeric chars with hyphens, collapse multiples."""
    slug = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
    return slug


def _parse_table_rows(lines: list[str]) -> list[list[str]]:
    """
    Parse markdown table lines into a list of rows (each row is a list of cell strings).
    Skips the separator row (e.g., |---|---|).
    """
    rows = []
    for line in lines:
        stripped = line.strip()
        if not stripped.startswith("|"):
            continue
        # Skip separator rows like |---|---|
        if re.match(r"^\|[\s\-:|]+\|$", stripped):
            continue
        cells = [c.strip() for c in stripped.split("|")[1:-1]]
        rows.append(cells)
    return rows


def _add_formatted_paragraph(doc: Document, text: str, style: str = "Normal"):
    """
    Add a paragraph with inline bold formatting.
    Handles **bold** markers within text.
    """
    para = doc.add_paragraph(style=style)
    # Split on **bold** markers
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = para.add_run(part[2:-2])
            run.bold = True
        else:
            para.add_run(part)
    return para


def write_contract_docx(contract_draft_md: str, company_name: str, reference: str) -> str:
    """
    Convert a markdown contract draft into a Word .docx file.

    Args:
        contract_draft_md: The markdown text of the contract.
        company_name: Company name (used in filename).
        reference: Reference number (e.g., OCX-atlantic-components-ltd-2026-05-11).

    Returns:
        Absolute path to the generated .docx file.
    """
    # Ensure contracts/ directory exists
    contracts_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "contracts")
    os.makedirs(contracts_dir, exist_ok=True)

    slug = _slugify(company_name)
    today = date.today().isoformat()
    filename = f"OCX-{slug}-{today}.docx"
    filepath = os.path.join(contracts_dir, filename)

    doc = Document()

    # ── Header ────────────────────────────────────────────────────────────────
    section = doc.sections[0]
    header = section.header
    header_para = header.paragraphs[0]
    header_para.text = "OceanX AI Ltd — Trade Finance Facility Agreement"
    header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ── Footer ────────────────────────────────────────────────────────────────
    footer = section.footer
    footer_para = footer.paragraphs[0]
    footer_para.text = f"{reference}"
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ── Parse markdown and build document ─────────────────────────────────────
    lines = contract_draft_md.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Headings
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
            i += 1
            continue
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
            i += 1
            continue
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
            i += 1
            continue

        # Detect lines that look like headings (all uppercase or short + ends with colon)
        if stripped and stripped.isupper() and len(stripped) < 60:
            doc.add_heading(stripped, level=2)
            i += 1
            continue
        if stripped and stripped.endswith(":") and len(stripped) < 60 and not stripped.startswith("-"):
            doc.add_heading(stripped[:-1], level=3)
            i += 1
            continue

        # Tables — collect consecutive lines starting with |
        if stripped.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            rows = _parse_table_rows(table_lines)
            if rows:
                num_cols = len(rows[0])
                table = doc.add_table(rows=len(rows), cols=num_cols)
                table.style = "Table Grid"
                for row_idx, row_data in enumerate(rows):
                    for col_idx, cell_text in enumerate(row_data):
                        if col_idx < num_cols:
                            table.rows[row_idx].cells[col_idx].text = cell_text
                # Bold the header row
                if rows:
                    for cell in table.rows[0].cells:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.bold = True
            continue

        # Bullet points
        if stripped.startswith("- ") or stripped.startswith("* "):
            bullet_text = stripped[2:]
            _add_formatted_paragraph(doc, bullet_text, style="List Bullet")
            i += 1
            continue

        # Empty lines
        if not stripped:
            i += 1
            continue

        # Normal paragraph with potential bold formatting
        _add_formatted_paragraph(doc, stripped)
        i += 1

    doc.save(filepath)
    return os.path.abspath(filepath)
